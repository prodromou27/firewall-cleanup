"""XLSX renderer — technical workbook: Summary + per-category finding sheets +
optional full rulebase / object inventory. Filterable headers, freeze panes,
colour-coded severity. Shares ReportData.
"""
import io

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from app.reporting.data import ReportData
from app.reporting import sections as S
from app.reporting.export_safety import spreadsheet_cell

_HEADER_FILL = PatternFill("solid", fgColor="1E3A5F")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_SEV_FILL = {
    "Critical": PatternFill("solid", fgColor="B91C1C"),
    "High": PatternFill("solid", fgColor="DC2626"),
    "Medium": PatternFill("solid", fgColor="D97706"),
    "Low": PatternFill("solid", fgColor="2563EB"),
    "Informational": PatternFill("solid", fgColor="6B7280"),
}


def _style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = _HEADER_FILL; cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _findings_sheet(ws, findings):
    cols = ["Severity", "Confidence", "Type", "Title", "Firewall", "Risk", "Affected rules", "Affected objects", "Evidence", "Recommendation"]
    ws.append(cols)
    for f in findings:
        ws.append([spreadsheet_cell(v) for v in [
            f["severity"], f["confidence"], f["finding_type_label"], f["title"],
            f["firewall_name"], f["risk_score"], f["affected_rule_count"],
            f["affected_object_count"], f["evidence_summary"], f["recommendation"],
        ]])
        ws.cell(row=ws.max_row, column=1).fill = _SEV_FILL.get(f["severity"], _SEV_FILL["Informational"])
        ws.cell(row=ws.max_row, column=1).font = Font(color="FFFFFF", bold=True)
    widths = [12, 11, 22, 46, 18, 7, 13, 14, 40, 60]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    _style_header(ws, len(cols))


def render(data: ReportData) -> bytes:
    wb = openpyxl.Workbook()
    m = data.meta

    # Summary
    ws = wb.active; ws.title = "Summary"
    ws.append([data.branding.get("report_title") or "Firewall Policy Review", ""])
    ws["A1"].font = Font(size=14, bold=True, color="1E3A5F")
    for k, v in [("Customer", m["customer_name"]), ("Firewall", m["firewall_name"]), ("Vendor", m["vendor"]),
                 ("Analysis date", m["analysis_date"]), ("Generated", m["generated_date"]),
                 ("Total rules", m["total_rules"]), ("Total objects", m["total_objects"]),
                 ("Total findings", m["total_findings"]), ("Critical", m["critical_findings"]),
                 ("High", m["high_findings"]), ("Medium", m["medium_findings"]),
                 ("Health score", m["policy_score"])]:
        ws.append([spreadsheet_cell(k), spreadsheet_cell(v)])
    ws.column_dimensions["A"].width = 20; ws.column_dimensions["B"].width = 40
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=1).font = Font(bold=True)

    # All findings
    if data.findings:
        _findings_sheet(wb.create_sheet("Findings"), data.findings)

    # Per-category finding sheets (from selected findings-type sections)
    seen = set()
    for s in data.sections:
        if s["type"] != S.T_FINDINGS or not s.get("findings"):
            continue
        title = s["name"][:31]
        if title in seen:
            continue
        seen.add(title)
        _findings_sheet(wb.create_sheet(title), s["findings"])

    # Appendices
    if data.rules and any(s["key"] == "full_rulebase" for s in data.sections):
        ws = wb.create_sheet("Full Rulebase")
        cols = ["#", "Name", "Source", "Destination", "Service", "Action", "Logging", "Hit count", "Last hit", "Risk"]
        ws.append(cols)
        for r in data.rules:
            ws.append([spreadsheet_cell(v) for v in [
                r["rule_number"], r["rule_name"], r["sources"], r["destinations"], r["services"],
                r["action"], r["logging"], r["hit_count"], r["last_hit"], r["risk_score"],
            ]])
        _style_header(ws, len(cols))
    if data.objects and any(s["key"] == "full_object_inventory" for s in data.sections):
        ws = wb.create_sheet("Object Inventory")
        cols = ["Name", "Type", "Value", "Members", "Member count"]
        ws.append(cols)
        for o in data.objects:
            ws.append([spreadsheet_cell(v) for v in [
                o["object_name"], o["object_type"], o["value"], o["members"], o["member_count"],
            ]])
        _style_header(ws, len(cols))

    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()
