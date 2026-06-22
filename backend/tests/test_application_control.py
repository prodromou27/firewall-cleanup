"""Tests for Check Point Application Control handling.

Covers the two false positives fixed for App-Control rules and the data-existence
gating the analysis relies on:

  * An Application Control rule (service "Any" but specific applications) is
    Layer-7 constrained and must NOT be reported as "overly permissive (any
    service)".
  * A genuinely wide-open rule (Any/Any/Any with no app constraint) is still
    flagged.
  * Zero-hit findings are only produced when hit-count data actually exists
    (_cp_hit_value returns None for "unknown", an int for real values).
  * Logging detection, disabled-rule handling, and application normalization for
    Check Point rules.
"""
import io

from app.analysis.engine import (
    _analyze_permissive, _analyze_usage, _analyze_disabled, _has_l7_app_constraint,
)
from app.connectors.live_sync import _cp_hit_value, _cp_logging


def _rule(**over):
    base = dict(
        id="r1", rule_id="1", rule_name="AppRule", enabled=True, action="accept",
        sources=["Internal-Net"], destinations=["Internal-Net"], services=["Any"],
        applications=[], hit_count=None,
    )
    base.update(over)
    return base


# ── False-positive fix: Application Control rules ─────────────────────────────
def test_app_control_rule_not_overly_permissive_on_service():
    """service=Any + specific applications → constrained at L7, no 'any service'."""
    r = _rule(applications=["Facebook", "YouTube-Streaming"])
    assert _analyze_permissive([r], {}) == []


def test_url_category_rule_not_overly_permissive():
    """A URL-filtering rule (content = category) is L7-constrained too."""
    r = _rule(applications=["Social-Networking", "Streaming-Media"])
    assert _analyze_permissive([r], {}) == []


def test_any_service_without_applications_still_flagged():
    r = _rule(applications=[])
    out = _analyze_permissive([r], {})
    assert out and "any service" in out[0]["title"]


def test_applications_any_is_not_a_constraint():
    """applications=['Any'] imposes no Layer-7 limit, so service Any still counts."""
    r = _rule(applications=["Any"])
    out = _analyze_permissive([r], {})
    assert out and "any service" in out[0]["title"]


def test_wide_open_rule_still_critical():
    r = _rule(sources=["Any"], destinations=["Any"], services=["Any"], applications=[])
    out = _analyze_permissive([r], {})
    assert out and out[0]["severity"] == "Critical"


def test_l7_constraint_helper():
    assert _has_l7_app_constraint(_rule(applications=["Facebook"]))
    assert not _has_l7_app_constraint(_rule(applications=[]))
    assert not _has_l7_app_constraint(_rule(applications=["Any"]))
    assert not _has_l7_app_constraint(_rule(applications=["", "any"]))


# ── Zero-hit gating: only when hit-count data exists ─────────────────────────
def test_cp_hit_value_unknown_when_no_data():
    assert _cp_hit_value({}) is None
    assert _cp_hit_value(None) is None


def test_cp_hit_value_real_values():
    assert _cp_hit_value({"value": 0}) == 0          # genuine zero
    assert _cp_hit_value({"value": 5}) == 5
    assert _cp_hit_value({"value": {"value": 3}}) == 3  # nested form


def test_zero_hit_finding_requires_hit_data():
    constrained = dict(services=["HTTP"], applications=[])
    # Unknown hits (None) → NO zero-hit finding (the false positive we fixed)
    none_rule = _rule(hit_count=None, **constrained)
    assert not [f for f in _analyze_usage([none_rule], {}) if f["finding_type"] == "zero_hit_rule"]
    # Real zero → flagged
    zero_rule = _rule(hit_count=0, **constrained)
    assert [f for f in _analyze_usage([zero_rule], {}) if f["finding_type"] == "zero_hit_rule"]
    # Positive hits → not flagged
    used_rule = _rule(hit_count=42, **constrained)
    assert not [f for f in _analyze_usage([used_rule], {}) if f["finding_type"] == "zero_hit_rule"]


def test_app_control_rule_no_false_zero_hit_when_hits_uncollected():
    """An App-Control rule synced without hit data must not produce a zero-hit finding."""
    r = _rule(applications=["Facebook"], services=["Any"], hit_count=_cp_hit_value({}))
    assert not [f for f in _analyze_usage([r], {}) if f["finding_type"] == "zero_hit_rule"]


# ── Logging detection (Check Point track field) ──────────────────────────────
def test_cp_logging_detection():
    assert _cp_logging({"type": {"name": "Log"}}) is True
    assert _cp_logging({"type": {"name": "Detailed Log"}}) is True
    assert _cp_logging({"type": {"name": "None"}}) is False
    assert _cp_logging({}) is False


# ── Disabled rules are excluded from active-rule findings ─────────────────────
def test_disabled_app_control_rule_reported_and_skipped_by_active_detectors():
    disabled = _rule(enabled=False, applications=["Facebook"], sources=["Any"], destinations=["Any"])
    # Disabled rules are surfaced by the disabled-rule detector…
    assert [f for f in _analyze_disabled([disabled], {}) if f["finding_type"] == "disabled_rule"]
    # …but not double-counted as overly permissive.
    assert _analyze_permissive([disabled], {}) == []


# ── Report exports surface the application column ────────────────────────────
def test_xlsx_full_rulebase_includes_applications():
    import openpyxl
    from app.reporting.data import ReportData
    from app.reporting.exporters import export

    rule = {
        "rule_number": 1, "rule_name": "Block social", "sources": "Internal",
        "destinations": "Any", "services": "Any", "applications": "Facebook, YouTube",
        "action": "accept", "logging": "Yes", "hit_count": "", "last_hit": "", "risk_score": 0,
    }
    data = ReportData(
        meta={"customer_name": "ACME", "firewall_name": "FW1", "vendor": "CheckPoint",
              "analysis_date": "", "generated_date": "", "total_rules": 1, "total_objects": 0,
              "total_findings": 0, "critical_findings": 0, "high_findings": 0, "medium_findings": 0,
              "policy_score": 90},
        branding={"report_title": "Tech"}, texts={},
        sections=[{"key": "full_rulebase", "name": "Full Rulebase", "type": "appendix"}],
        findings=[], severity_counts={}, category_counts={}, rules=[rule], objects=[],
    )
    content, _, ext = export(data, "xlsx")
    assert ext == "xlsx"
    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb["Full Rulebase"]
    header = [c.value for c in ws[1]]
    assert "Application" in header
    assert any("Facebook" in str(c.value) for row in ws.iter_rows() for c in row)
